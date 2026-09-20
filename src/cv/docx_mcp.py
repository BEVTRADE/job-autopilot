"""
Édition des CV Word par le serveur MCP docx-mcp-server (SecurityRonin).

Pourquoi un serveur MCP plutôt que render.substitute : le remplacement
littéral dans le XML échoue en silence quand le texte est découpé en
plusieurs segments (runs). C'est le cas du titre des CV maîtres, découpé en
19 segments : la substitution n'y change rien et ne le signale pas. Le
serveur, lui, travaille au niveau du paragraphe et conserve la mise en forme.

Constats du 19/09/2026, vérifiés sur FR_ARCHITECTE_IA_GENAI.docx :
- docx-mcp-server 0.7.4 exige mcp<2 (FastMCP renommé en mcp 2.x) ;
- il cible les paragraphes par w14:paraId, absent des CV générés par
  python-docx : on les ajoute d'abord (ajouter_para_ids) ;
- il expose 219 outils, environ 29 000 jetons de description : jamais
  présentés tels quels à un modèle local (voir OUTILS_AGENT).

Le modèle ne pilote pas ce serveur : le code déterministe l'appelle avec des
remplacements déjà validés. Le modèle ne touche jamais au fichier.
"""
from __future__ import annotations
import asyncio, json, os, random, re, shutil, zipfile

W14 = 'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"'
# Sous-ensemble à exposer si un agent interactif doit un jour éditer un CV.
OUTILS_AGENT = ("open_document", "get_body_text", "search_text",
                "replace_text", "save_document", "close_document")


def commande_serveur() -> tuple[str, list[str]]:
    """Le serveur installé dans le venv, sinon via uvx avec mcp épinglé."""
    import sys
    local = os.path.join(os.path.dirname(sys.executable), "docx-mcp-server")
    if os.path.exists(local):
        return local, []
    return "uvx", ["--with", "mcp<2", "--from", "docx-mcp-server", "docx-mcp-server"]


def ajouter_para_ids(src: str, dst: str) -> int:
    """Copie src vers dst en donnant un w14:paraId unique à chaque paragraphe."""
    n = 0
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for it in zin.infolist():
            d = zin.read(it.filename)
            if it.filename == "word/document.xml":
                t = d.decode("utf-8")
                if "xmlns:w14=" not in t:
                    t = t.replace("<w:document ", f"<w:document {W14} ", 1)
                pris = set(re.findall(r'w14:paraId="([0-9A-F]{8})"', t))

                def f(m):
                    nonlocal n
                    if "w14:paraId" in m.group(0):
                        return m.group(0)
                    while True:
                        pid = "%08X" % random.randint(1, 0x7FFFFFFF)
                        if pid not in pris:
                            pris.add(pid)
                            break
                    n += 1
                    return m.group(0)[:-1] + f' w14:paraId="{pid}" w14:textId="77777777">'
                t = re.sub(r"<w:p(?=[ >])[^>]*>", f, t)
                d = t.encode("utf-8")
            z.writestr(it, d)
    return n


class EchecEdition(RuntimeError):
    pass


async def _appliquer(src: str, dst: str, substitutions: list[dict]) -> list[dict]:
    from mcp import ClientSession, StdioServerParameters
    from mcp.client.stdio import stdio_client
    cmd, args = commande_serveur()
    bilan = []
    async with stdio_client(StdioServerParameters(command=cmd, args=args)) as (r, w):
        async with ClientSession(r, w) as s:
            await s.initialize()
            o = await s.call_tool("open_document", {"path": os.path.abspath(src)})
            if o.isError:
                raise EchecEdition(f"ouverture refusée : {o.content[0].text[:200]}")
            for sub in substitutions:
                av, ap = sub["avant"], sub["apres"]
                t = await s.call_tool("search_text", {"query": av})
                trouves = [x for x in json.loads(t.content[0].text) if x.get("paraId")] if not t.isError else []
                if len(trouves) != 1:
                    bilan.append({**sub, "applique": False,
                                  "motif": "texte introuvable" if not trouves else f"{len(trouves)} occurrences"})
                    continue
                x = await s.call_tool("replace_text", {"para_id": trouves[0]["paraId"], "find": av,
                                                       "replace": ap, "tracked": False})
                bilan.append({**sub, "applique": not x.isError,
                              "motif": None if not x.isError else x.content[0].text[:200]})
            sv = await s.call_tool("save_document", {"output_path": os.path.abspath(dst)})
            if sv.isError:
                raise EchecEdition(f"enregistrement refusé : {sv.content[0].text[:200]}")
    return bilan


W14_NS = "http://schemas.microsoft.com/office/word/2010/wordml"


async def _remplacer(prep: str, dst: str, remplacements: list[dict]) -> list[dict]:
    """Remplace le texte de paragraphes désignés par leur paraId, sans suivi des modifications."""
    from mcp import ClientSession, StdioServerParameters
    from mcp.client.stdio import stdio_client
    cmd, args = commande_serveur()
    bilan = []
    async with stdio_client(StdioServerParameters(command=cmd, args=args)) as (r, w):
        async with ClientSession(r, w) as s:
            await s.initialize()
            o = await s.call_tool("open_document", {"path": os.path.abspath(prep)})
            if o.isError:
                raise EchecEdition(f"ouverture refusée : {o.content[0].text[:200]}")
            for x in remplacements:
                if x.get("deplacement"):
                    # Paragraphe désigné par son paraId, sans recherche de texte : replace_text refuse
                    # « Ambiguous » dès qu'un texte déplacé existe deux fois dans le document.
                    y = await s.call_tool("update_paragraph", {"para_id": x["para_id"], "text": x["apres"]})
                else:
                    y = await s.call_tool("replace_text", {"para_id": x["para_id"], "find": x["avant"],
                                                           "replace": x["apres"], "tracked": False})
                bilan.append({"indice": x["indice"], "applique": not y.isError,
                              "motif": None if not y.isError else y.content[0].text[:200]})
            sv = await s.call_tool("save_document", {"output_path": os.path.abspath(dst)})
            if sv.isError:
                raise EchecEdition(f"enregistrement refusé : {sv.content[0].text[:200]}")
    return bilan


def appliquer_remplacements(src: str, dst: str, remplacements: list[dict]) -> list[dict]:
    """Remplace, dans le paragraphe d'indice donné (python-docx, à partir de 0), le texte `avant` par `apres`.
    Un remplacement marqué `deplacement` réécrit le paragraphe entier (texte d'un autre paragraphe du CV).

    Tout ou rien : si un remplacement échoue ou si le fichier relu ne contient pas exactement les textes
    attendus, EchecEdition est levée et `dst` n'existe pas (un ordre d'accroche à moitié appliqué
    dupliquerait ou perdrait un paragraphe).
    """
    from docx import Document
    os.makedirs(os.path.dirname(os.path.abspath(dst)), exist_ok=True)
    prep, brut = dst + ".ids.docx", dst + ".tmp.docx"
    ajouter_para_ids(src, prep)
    try:
        paras = Document(prep).paragraphs
        for x in remplacements:
            x["para_id"] = paras[x["indice"]]._p.get(f"{{{W14_NS}}}paraId")
            if not x["para_id"] or paras[x["indice"]].text != x["avant"]:
                raise EchecEdition(f"paragraphe {x['indice']} absent ou différent de celui attendu")
        bilan = asyncio.run(_remplacer(prep, brut, remplacements)) if remplacements else []
        if remplacements:
            echecs = [b for b in bilan if not b["applique"]]
            if echecs:
                raise EchecEdition(f"remplacement refusé : {echecs[0]['motif']}")
            relu = Document(brut).paragraphs
            attendu = {x["indice"]: x["apres"] for x in remplacements}
            if len(relu) != len(paras) or any(relu[i].text != t for i, t in attendu.items()):
                raise EchecEdition("le fichier relu ne contient pas les textes attendus")
            os.replace(brut, dst)
        else:
            shutil.copy(prep, dst)
        return bilan
    finally:
        for f in (prep, brut):
            if os.path.exists(f):
                try:
                    os.remove(f)
                except OSError:          # montage sans droit de suppression
                    pass


def appliquer(src: str, dst: str, substitutions: list[dict]) -> list[dict]:
    """Applique des substitutions validées ; renvoie le bilan, une ligne par substitution."""
    os.makedirs(os.path.dirname(os.path.abspath(dst)), exist_ok=True)
    prep = dst + ".ids.docx"
    ajouter_para_ids(src, prep)
    try:
        if not substitutions:
            shutil.copy(prep, dst)
            return []
        return asyncio.run(_appliquer(prep, dst, substitutions))
    finally:
        if os.path.exists(prep):
            try:
                os.remove(prep)
            except OSError:          # montage sans droit de suppression
                pass
